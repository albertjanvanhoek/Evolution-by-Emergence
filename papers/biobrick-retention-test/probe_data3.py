"""
Probe the published Supplementary Data 3 workbook from:
Radde et al., Nature Communications 15, 6242 (2024)
DOI: 10.1038/s41467-024-50639-9

This script only inspects workbook structure so the exact published
301-BioBrick cohort can be reconstructed reproducibly.
"""
from urllib.request import Request, urlopen
from pathlib import Path
from openpyxl import load_workbook

URL = "https://media.springernature.com/original/springer-static/esm/art%3A10.1038%2Fs41467-024-50639-9/MediaObjects/41467_2024_50639_MOESM6_ESM.xlsx"
OUT = Path("/tmp/biobrick_data3.xlsx")

req = Request(URL, headers={"User-Agent": "Mozilla/5.0"})
with urlopen(req, timeout=60) as r:
    OUT.write_bytes(r.read())

print("downloaded", OUT.stat().st_size, "bytes")
wb = load_workbook(OUT, read_only=True, data_only=True)
print("sheets:", wb.sheetnames)
for ws in wb.worksheets:
    print("\nSHEET", ws.title, "rows", ws.max_row, "cols", ws.max_column)
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        nonnull = sum(v is not None for v in row)
        if i <= 40 or nonnull >= 5:
            print("ROW", i, "nonnull", nonnull, repr(row))
        if i >= 40:
            break
