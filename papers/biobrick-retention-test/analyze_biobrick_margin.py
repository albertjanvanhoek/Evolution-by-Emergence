"""
Exact-cohort reanalysis of the published BioBrick burden data.

Source:
Radde et al. (2024), Nature Communications 15:6242
DOI 10.1038/s41467-024-50639-9
Supplementary Data 3.

Goals:
1. Reconstruct the exact published 301-BioBrick cohort (excluding 5 BFP controls).
2. Transform fractional burden b to the uniform-dilution load kappa=b/(1-b).
3. Explore whether the upper tail of kappa is consistent with a finite endpoint.
4. Identify candidate constructs for a prospective two-load test.

This is an exploratory bridge analysis, not evidence that the inferred endpoint
is an intrinsic E. coli constant. The observed cohort is already subject to
constructability/evolutionary selection.
"""
from __future__ import annotations

from urllib.request import Request, urlopen
from pathlib import Path
from math import isfinite
import itertools
import json

import numpy as np
from openpyxl import load_workbook
from scipy.stats import genpareto

URL = "https://media.springernature.com/original/springer-static/esm/art%3A10.1038%2Fs41467-024-50639-9/MediaObjects/41467_2024_50639_MOESM6_ESM.xlsx"
CACHE = Path("/tmp/biobrick_data3.xlsx")
SEED = 20260914
BOOT = 1000
BURDEN_THRESHOLDS = [0.12, 0.15, 0.18, 0.20, 0.22, 0.25, 0.26]


def download() -> Path:
    req = Request(URL, headers={"User-Agent": "Mozilla/5.0"})
    with urlopen(req, timeout=60) as r:
        CACHE.write_bytes(r.read())
    return CACHE


def load_rows(path: Path):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Table S3 - BioBrick part result"]
    header_row = None
    header = None
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row and row[0] == "accession":
            header_row = i
            header = list(row)
            break
    if header_row is None:
        raise RuntimeError("Could not locate Supplementary Data 3 header")
    rows = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if row[0] is None:
            continue
        rows.append(dict(zip(header, row)))
    return rows


def finite_float(x):
    try:
        y = float(x)
        return y if isfinite(y) else None
    except (TypeError, ValueError):
        return None


def prepare(rows):
    controls, parts = [], []
    for r in rows:
        g = finite_float(r["normalized.growth.rate.mean"])
        if g is None:
            continue
        b = 1.0 - g
        kappa = b / (1.0 - b)
        out = dict(r)
        out["burden"] = b
        out["kappa"] = kappa
        if str(r["burden.category"]).strip().lower() == "control":
            controls.append(out)
        else:
            parts.append(out)
    return controls, parts


def fit_gpd(values, u):
    exc = np.asarray([x - u for x in values if x > u], dtype=float)
    if len(exc) < 10:
        return None
    xi, loc, scale = genpareto.fit(exc, floc=0.0)
    endpoint = np.inf if xi >= 0 else u - scale / xi
    return {
        "n": int(len(exc)),
        "xi": float(xi),
        "scale": float(scale),
        "endpoint_kappa": float(endpoint),
        "endpoint_burden": float(endpoint / (1.0 + endpoint)) if np.isfinite(endpoint) else np.inf,
    }


def bootstrap(values, u, B=BOOT):
    rng = np.random.default_rng(SEED + int(round(1000 * u)))
    vals = np.asarray(values, dtype=float)
    xis, ends = [], []
    for _ in range(B):
        sample = rng.choice(vals, size=len(vals), replace=True)
        fit = fit_gpd(sample, u)
        if fit is None:
            continue
        xis.append(fit["xi"])
        if fit["xi"] < 0 and np.isfinite(fit["endpoint_kappa"]):
            ends.append(fit["endpoint_kappa"])

    def q(a, p):
        return float(np.quantile(np.asarray(a), p)) if a else np.nan

    return {
        "valid": len(xis),
        "negative_fraction": float(np.mean(np.asarray(xis) < 0)) if xis else np.nan,
        "xi_ci": [q(xis, 0.025), q(xis, 0.975)],
        "endpoint_kappa_ci_cond_negative": [q(ends, 0.025), q(ends, 0.975)],
    }


def candidate_constructs(parts):
    out = []
    for r in parts:
        b = r["burden"]
        if not (0.18 <= b <= 0.32):
            continue
        if str(r["burden.category"]).strip().lower() != "significant":
            continue
        if bool(r["GFP.interference"]):
            continue
        other = str(r["other.burden.greater.significant"]).strip().lower()
        if other == "significant":
            continue
        out.append(r)
    return out


def candidate_pairs(cands):
    pairs = []
    for a, b in itertools.combinations(cands, 2):
        b1, b2 = a["burden"], b["burden"]
        add = b1 + b2
        mult = 1.0 - (1.0 - b1) * (1.0 - b2)
        if add >= 0.48 and mult <= 0.46:
            score = abs(add - 0.50) + abs(mult - 0.44)
            pairs.append({
                "a": a["accession"],
                "b": b["accession"],
                "b1": b1,
                "b2": b2,
                "additive": add,
                "multiplicative": mult,
                "score": score,
            })
    return sorted(pairs, key=lambda x: x["score"])


def main():
    rows = load_rows(download())
    controls, parts = prepare(rows)

    print("PUBLISHED COHORT")
    print("all rows:", len(rows))
    print("BFP controls:", len(controls))
    print("BioBricks:", len(parts))
    assert len(rows) == 306, len(rows)
    assert len(controls) == 5, len(controls)
    assert len(parts) == 301, len(parts)

    burdens = np.asarray([r["burden"] for r in parts])
    kappas = np.asarray([r["kappa"] for r in parts])

    print("burden min/max:", float(burdens.min()), float(burdens.max()))
    print("kappa min/max:", float(kappas.min()), float(kappas.max()))
    print("BioBricks >45% burden:", int(np.sum(burdens > 0.45)))
    print("BioBricks >40% burden:", int(np.sum(burdens > 0.40)))

    print("\nBFP CONTROLS")
    for r in sorted(controls, key=lambda x: x["burden"]):
        print(r["accession"], r["strains"], f"b={r['burden']:.6f}", f"kappa={r['kappa']:.6f}")

    print("\nGPD ON KAPPA (thresholds defined on burden scale)")
    results = []
    for ub in BURDEN_THRESHOLDS:
        uk = ub / (1.0 - ub)
        fit = fit_gpd(kappas, uk)
        if fit is None:
            continue
        boot = bootstrap(kappas, uk)
        rec = {"threshold_burden": ub, "threshold_kappa": uk, **fit, **boot}
        results.append(rec)
        print(
            f"b>{ub:.2f} n={fit['n']:2d} "
            f"xi={fit['xi']:+.3f} "
            f"M_end={fit['endpoint_kappa']:.3f} "
            f"b_end={fit['endpoint_burden']:.3f} "
            f"boot P(xi<0)={boot['negative_fraction']:.3f} "
            f"xi95=[{boot['xi_ci'][0]:+.3f},{boot['xi_ci'][1]:+.3f}] "
            f"M95|xi<0=[{boot['endpoint_kappa_ci_cond_negative'][0]:.3f},"
            f"{boot['endpoint_kappa_ci_cond_negative'][1]:.3f}]"
        )

    print("\nTRANSLATION-DOMINATED CANDIDATES (screen only)")
    cands = candidate_constructs(parts)
    print("candidate count:", len(cands))
    for r in sorted(cands, key=lambda x: -x["burden"])[:20]:
        print(
            r["accession"],
            f"b={r['burden']:.4f}",
            f"kappa={r['kappa']:.4f}",
            "vector=", r["vectors"],
            "other=", r["other.burden.greater.significant"],
        )

    print("\nDISCRIMINATING PAIRS")
    pairs = candidate_pairs(cands)
    print("pair count:", len(pairs))
    for p in pairs[:15]:
        print(
            p["a"], p["b"],
            f"b1={p['b1']:.4f}", f"b2={p['b2']:.4f}",
            f"add={p['additive']:.4f}",
            f"mult={p['multiplicative']:.4f}",
        )

    summary = {
        "source_doi": "10.1038/s41467-024-50639-9",
        "n_biobricks": len(parts),
        "n_controls": len(controls),
        "max_burden": float(burdens.max()),
        "max_kappa": float(kappas.max()),
        "gpd_kappa": results,
        "candidate_count": len(cands),
        "discriminating_pair_count": len(pairs),
        "top_pairs": pairs[:15],
    }
    Path("/tmp/biobrick_retention_summary.json").write_text(json.dumps(summary, indent=2))
    print("\nALL EXACT-COHORT CHECKS PASS")


if __name__ == "__main__":
    main()
